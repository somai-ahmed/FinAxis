"""Export Excel (openpyxl) des rapports FinCore."""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook

from .templates import balance_template, bilan_template, detection_template, grand_livre_template


def _sauver(wb: Workbook, chemin: str) -> str:
    if not chemin.lower().endswith(".xlsx"):
        chemin += ".xlsx"
    wb.save(chemin)
    return chemin


def exporter_balance(chemin, lignes, entreprise, debut: date, fin: date, devise="TND") -> str:
    wb = Workbook()
    balance_template.ecrire(wb.active, lignes, entreprise, debut, fin, devise)
    return _sauver(wb, chemin)


def exporter_grand_livre(chemin, lignes, compte_libelle, entreprise, debut: date, fin: date, devise="TND") -> str:
    wb = Workbook()
    grand_livre_template.ecrire(wb.active, lignes, compte_libelle, entreprise, debut, fin, devise)
    return _sauver(wb, chemin)


def exporter_bilan(chemin, bilan, entreprise, date_arret: date, devise="TND") -> str:
    wb = Workbook()
    bilan_template.ecrire(wb.active, bilan, entreprise, date_arret, devise)
    return _sauver(wb, chemin)


def exporter_detection(chemin, resultats, comptes_par_id, entreprise) -> str:
    wb = Workbook()
    detection_template.ecrire(wb.active, resultats, comptes_par_id, entreprise)
    return _sauver(wb, chemin)


def exporter_classeur_complet(chemin, *, entreprise, debut: date, fin: date, devise, balance, bilan, resultats, comptes_par_id) -> str:
    """Un seul classeur : Balance + Bilan + Detection."""
    wb = Workbook()
    balance_template.ecrire(wb.active, balance, entreprise, debut, fin, devise)
    bilan_template.ecrire(wb.create_sheet(), bilan, entreprise, fin, devise)
    detection_template.ecrire(wb.create_sheet(), resultats, comptes_par_id, entreprise)
    return _sauver(wb, chemin)
