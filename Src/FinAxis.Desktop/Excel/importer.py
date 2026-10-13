"""Import Excel d'ecritures (stretch).

Format attendu, une ligne d'en-tete puis une ligne par ligne d'ecriture :
    Date | Référence | Libellé | Code compte | Débit | Crédit
Les lignes qui partagent la meme Référence forment UNE ecriture.
"""
from __future__ import annotations

from datetime import date, datetime

from openpyxl import load_workbook

from helpers.currency_helper import analyser_montant, decimal_vers_millimes
from helpers.date_helper import analyser_date
from models.journal_entry import JournalEntry, JournalLine


class ErreurImport(Exception):
    def __init__(self, ligne: int, message: str):
        super().__init__(f"Ligne {ligne} : {message}")
        self.ligne = ligne


def _en_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return analyser_date(str(v)) if v is not None else None


def _en_millimes(v) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return decimal_vers_millimes(v)
    return analyser_montant(str(v)) or 0


def importer_ecritures(chemin: str, id_par_code: dict[str, int]) -> list[JournalEntry]:
    """Lit le classeur et retourne les ecritures (brouillons). Leve ErreurImport au premier probleme."""
    feuille = load_workbook(chemin, read_only=True, data_only=True).active
    groupes: dict[str, JournalEntry] = {}
    for numero, ligne in enumerate(feuille.iter_rows(min_row=2, values_only=True), start=2):
        if not ligne or all(c is None for c in ligne):
            continue
        d, ref, libelle, code, debit, credit = (list(ligne) + [None] * 6)[:6]
        d = _en_date(d)
        if d is None:
            raise ErreurImport(numero, "date invalide")
        ref = str(ref or "").strip()
        if not ref:
            raise ErreurImport(numero, "référence manquante")
        code = str(code or "").strip()
        if code not in id_par_code:
            raise ErreurImport(numero, f"compte inconnu « {code} »")
        ecriture = groupes.setdefault(ref, JournalEntry(d, ref, str(libelle or ref)))
        ecriture.lignes.append(JournalLine(id_par_code[code], _en_millimes(debit), _en_millimes(credit), str(libelle or "")))
    return list(groupes.values())
