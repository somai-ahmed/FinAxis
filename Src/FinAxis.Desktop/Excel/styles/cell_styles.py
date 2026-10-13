"""Styles reutilisables (openpyxl)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import color_palette as c

FORMAT_MONTANT = '#,##0.000;[Red]-#,##0.000'
FORMAT_DATE = "dd/mm/yyyy"

_fin = Side(style="thin", color="BFBFBF")
BORDURE = Border(left=_fin, right=_fin, top=_fin, bottom=_fin)

TITRE = Font(name="Calibri", size=14, bold=True, color=c.BLEU_FONCE)
SOUS_TITRE = Font(name="Calibri", size=10, italic=True, color="595959")
ENTETE_FONT = Font(name="Calibri", bold=True, color=c.BLANC)
ENTETE_FOND = PatternFill("solid", fgColor=c.BLEU_FONCE)
TOTAL_FONT = Font(name="Calibri", bold=True)
TOTAL_FOND = PatternFill("solid", fgColor=c.BLEU_CLAIR)
CENTRE = Alignment(horizontal="center", vertical="center")
DROITE = Alignment(horizontal="right")


def fond(couleur: str) -> PatternFill:
    return PatternFill("solid", fgColor=couleur)


def ecrire_entetes(ws, ligne: int, entetes: list[str]) -> None:
    for col, texte in enumerate(entetes, start=1):
        cell = ws.cell(row=ligne, column=col, value=texte)
        cell.font, cell.fill, cell.alignment, cell.border = ENTETE_FONT, ENTETE_FOND, CENTRE, BORDURE


def largeurs(ws, valeurs: list[int]) -> None:
    from openpyxl.utils import get_column_letter
    for i, v in enumerate(valeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = v
