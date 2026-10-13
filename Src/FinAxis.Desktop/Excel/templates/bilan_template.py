"""Feuille 'Bilan' (15 lignes, avec indentation selon la profondeur)."""
from datetime import date

from openpyxl.styles import Alignment

from ..styles import cell_styles as st


def ecrire(ws, bilan, entreprise: str, date_arret: date, devise: str = "TND") -> None:
    ws.title = "Bilan"
    ws["A1"], ws["A1"].font = f"Bilan - {entreprise}", st.TITRE
    ws["A2"], ws["A2"].font = f"Arrêté au {date_arret:%d/%m/%Y} ({devise})", st.SOUS_TITRE
    st.ecrire_entetes(ws, 4, ["Poste", "Montant"])
    ligne = 5
    for r in bilan.lignes:
        a = ws.cell(row=ligne, column=1, value=r.libelle)
        b = ws.cell(row=ligne, column=2, value=r.montant / 1000)
        a.alignment = Alignment(indent=2 * r.profondeur)
        b.number_format = st.FORMAT_MONTANT
        for cell in (a, b):
            cell.border = st.BORDURE
            if r.est_sous_total:
                cell.font, cell.fill = st.TOTAL_FONT, st.TOTAL_FOND
        ligne += 1
    ws.cell(row=ligne + 1, column=1, value="Équilibré" if bilan.est_equilibre else "ATTENTION : actif ≠ passif")
    ws.cell(row=ligne + 1, column=1).font = st.TOTAL_FONT
    ws.freeze_panes = "A5"
    st.largeurs(ws, [48, 20])
